"""生成各格式测试样例，用于自动化回归。"""
from __future__ import annotations

import os
import zipfile
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "samples"
OUT.mkdir(exist_ok=True)

ZH = [
    "这是一份关于软件采购与技术支持服务的合同。",
    "供应商应在 2026-12-31 前完成数据库迁移与接口联调。",
    "付款条件：验收合格后 30 日内支付合同总额的 80%。",
    "违约责任：任何一方违反保密义务，应承担相应的损害赔偿责任。",
]


def make_docx() -> Path:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("技术服务合同", level=1)
    doc.add_paragraph(ZH[0])
    p = doc.add_paragraph()
    p.add_run("甲方：").bold = True
    p.add_run("示例科技有限公司")
    doc.add_paragraph(ZH[1])
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "阶段"
    table.cell(0, 1).text = "交付物"
    table.cell(1, 0).text = "第一阶段"
    table.cell(1, 1).text = ZH[2]
    doc.add_paragraph(ZH[3])
    doc.sections[0].header.paragraphs[0].text = "机密文件 · 请勿外传"
    path = OUT / "sample.docx"
    doc.save(path)
    return path


def make_xlsx() -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "报价单"
    ws["A1"] = "项目名称"
    ws["B1"] = "数量"
    ws["C1"] = "单价"
    ws["D1"] = "备注"
    ws["A2"] = ZH[0]
    ws["B2"] = 10
    ws["C2"] = 1200
    ws["D2"] = ZH[2]
    ws["A3"] = ZH[1]
    ws["B3"] = 2
    ws["C3"] = 5000
    ws["D3"] = "含税"
    ws["E2"] = "=B2*C2"
    ws["E3"] = "=B3*C3"
    ws.merge_cells("A5:D5")
    ws["A5"] = ZH[3]
    ws.column_dimensions["A"].width = 40
    path = OUT / "sample.xlsx"
    wb.save(path)
    return path


def make_pptx() -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "产品发布计划"
    body = slide.placeholders[1]
    body.text_frame.text = ZH[0]
    body.text_frame.add_paragraph().text = ZH[1]
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "里程碑"
    rows, cols = 2, 2
    table = slide2.shapes.add_table(rows, cols, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    table.cell(0, 0).text = "阶段"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "研发"
    table.cell(1, 1).text = ZH[2]
    path = OUT / "sample.pptx"
    prs.save(path)
    return path


def make_md_txt() -> tuple[Path, Path]:
    md = (
        "---\n"
        "title: 使用手册\n"
        "---\n\n"
        "# 快速开始\n\n"
        "本工具支持批量翻译，详见 [文档](https://example.com/docs)。\n\n"
        "```python\n"
        "print('这段代码不应被翻译')\n"
        "```\n\n"
        "- 第一步：选择文件\n"
        "- 第二步：设置术语库 `terms.json`\n\n"
        "> 注意：`API` Key 请勿外传。\n"
    )
    p1 = OUT / "sample.md"
    p1.write_text(md, "utf-8")

    txt = (
        "这是一段被硬换行拆开的\n"
        "长文本，需要先合并再翻译。\n\n"
        "第二个段落包含网址 https://example.com 与邮箱 a@b.com。\n"
    )
    p2 = OUT / "sample.txt"
    p2.write_text(txt, "utf-8")
    return p1, p2


VSDX_PAGE = """<?xml version='1.0' encoding='utf-8' standalone='yes'?>
<PageContents xmlns="http://schemas.microsoft.com/office/visio/2012/main">
  <Shapes>
    <Shape ID="1" Type="Shape" LineStyle="0" FillStyle="0" TextStyle="0">
      <Cell N="PinX" V="4"/><Cell N="PinY" V="6"/><Cell N="Width" V="3"/><Cell N="Height" V="1"/>
      <Text>审批流程说明</Text>
    </Shape>
    <Shape ID="2" Type="Shape">
      <Cell N="PinX" V="4"/><Cell N="PinY" V="4"/><Cell N="Width" V="3"/><Cell N="Height" V="1"/>
      <Text><cp IX='0'/>由部门负责人<cp IX='1'/>签字确认</Text>
    </Shape>
  </Shapes>
</PageContents>
"""


def make_vsdx() -> Path:
    path = OUT / "sample.vsdx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        z.writestr("visio/pages/pages.xml",
                   '<?xml version="1.0" encoding="utf-8"?><Pages xmlns="http://schemas.microsoft.com/office/visio/2012/main"/>')
        z.writestr("visio/pages/page1.xml", VSDX_PAGE)
    return path


XHTML = """<?xml version="1.0" encoding="utf-8"?>
<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Chapter 1</title></head>
<body>
<h1>第一章 引言</h1>
<p>{p0}</p>
<p>参见 <a href="https://example.com">官方网站</a> 获取更多信息。</p>
<p>{p1}</p>
</body></html>
"""


def make_epub() -> Path:
    path = OUT / "sample.epub"
    text = XHTML.format(p0=ZH[0], p1=ZH[2])
    opf = """<?xml version="1.0" encoding="utf-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id">
<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>示例书</dc:title></metadata>
<manifest><item id="c1" href="chapter1.xhtml" media-type="application/xhtml+xml"/></manifest>
<spine><itemref idref="c1"/></spine>
</package>"""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                   '<rootfiles><rootfile full-path="content.opf" media-type="application/oebps-package+xml"/>'
                   '</rootfiles></container>', compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("content.opf", opf, compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("chapter1.xhtml", text, compress_type=zipfile.ZIP_DEFLATED)
    return path




def make_pdf() -> Path:
    import fitz

    doc = fitz.open()
    page = doc.new_page(width=460, height=320)
    page.insert_text((46, 70), "Technical Service Agreement", fontsize=15)
    page.insert_text((46, 104), ZH[0], fontsize=10, fontname="china-s")
    page.draw_rect(fitz.Rect(40, 130, 420, 250), color=(0.1, 0.2, 0.6), width=1)
    page.insert_text((56, 158), "Party A: Sample Technology Co., Ltd.", fontsize=10)
    page.insert_text((56, 186), "Deliverables: database migration and API integration.", fontsize=10)
    page.insert_text((56, 214), "Payment terms: 30 days net after acceptance.", fontsize=10)
    page.insert_text((46, 286), "Confidential - do not distribute.", fontsize=9)
    path = OUT / "sample.pdf"
    doc.save(path)
    doc.close()
    return path


def main() -> None:
    made = [make_docx(), make_xlsx(), make_pptx(), make_pdf(), *make_md_txt(), make_vsdx(), make_epub()]
    for p in made:
        print(f"{p.name:16s} {p.stat().st_size:>8d} bytes")


if __name__ == "__main__":
    main()
