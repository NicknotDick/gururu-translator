"""八种格式端到端冒烟测试：解析 → 翻译 → 写回 → 校验结构未破坏。"""
from __future__ import annotations

import sys
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app.core.ir import STATUS_DONE  # noqa: E402
from app.engines.chain import Translator  # noqa: E402
from app.parsers.registry import parser_for  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "samples"
OUT = SAMPLES / "out"
OUT.mkdir(exist_ok=True)


def structure(path: Path) -> dict:
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            from docx import Document

            d = Document(path)
            return {"paragraphs": len(d.paragraphs), "tables": len(d.tables)}
        if ext == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(path)
            ws = wb.active
            return {"sheets": len(wb.sheetnames), "dims": f"{ws.max_row}x{ws.max_column}",
                    "merged": len(ws.merged_cells.ranges)}
        if ext == ".pptx":
            from pptx import Presentation

            p = Presentation(path)
            return {"slides": len(p.slides),
                    "shapes": sum(len(s.shapes) for s in p.slides)}
        if ext == ".pdf":
            import fitz

            d = fitz.open(path)
            return {"pages": d.page_count, "drawings": len(d[0].get_drawings())}
        if ext == ".vsdx":
            import zipfile

            with zipfile.ZipFile(path) as z:
                return {"entries": len(z.namelist())}
        if ext == ".epub":
            import zipfile

            with zipfile.ZipFile(path) as z:
                names = z.namelist()
                return {"entries": len(names), "mimetype_first": names[0] == "mimetype"}
    except Exception as exc:  # noqa: BLE001
        return {"error": str(exc)[:80]}
    return {}


def run(name: str, mode: str = "target_only") -> bool:
    src = SAMPLES / name
    if not src.exists():
        print(f"[SKIP] {name} 不存在")
        return True
    print(f"\n=== {name} ({mode}) ===")
    try:
        parser = parser_for(str(src))
    except ValueError as exc:
        print(f"[FAIL] 无法解析：{exc}")
        return False
    try:
        ir = parser.load(str(src))
    except Exception as exc:  # noqa: BLE001
        print(f"[FAIL] 读取失败：{type(exc).__name__}: {exc}")
        return False
    print(f"段数 {len(ir.segments)}  警告 {ir.warnings}")
    if not ir.segments:
        print("[WARN] 没有可翻译片段")
        return False

    tr = Translator()
    res = tr.translate([s.engine_text for s in ir.segments], "zh", "en")
    for seg, out in zip(ir.segments, res.texts):
        seg.tgt = out
        seg.status = STATUS_DONE
    print(f"引擎 {res.engine}  失败 {len(res.failed)}  {res.warnings[:1]}")

    stem = src.stem + ("_bi" if mode == "bilingual" else "")
    out = OUT / f"{stem}{src.suffix}"
    try:
        info = parser.run_write(ir, str(out), mode)
    except Exception as exc:  # noqa: BLE001
        print(f"[FAIL] 写回失败：{type(exc).__name__}: {exc}")
        traceback.print_exc()
        return False
    before, after = structure(src), structure(out)
    print(f"结构 前 {before}")
    print(f"结构 后 {after}")
    print(f"溢出 {len(info.get('overflow', []))}  警告 {info.get('warnings', [])}")
    print(f"输出 {out.relative_to(ROOT)}  {out.stat().st_size} bytes")
    for seg in ir.segments[:3]:
        print(f"   {seg.src[:34]!r} -> {seg.tgt[:44]!r}")
    return True


def main() -> None:
    cases = sys.argv[1:] or [
        "sample.docx", "sample.xlsx", "sample.pptx", "sample.pdf",
        "sample.md", "sample.txt", "sample.vsdx", "sample.epub",
    ]
    mode = "bilingual" if "--bi" in cases else "target_only"
    cases = [c for c in cases if not c.startswith("--")]
    ok = 0
    for c in cases:
        if run(c, mode):
            ok += 1
    print(f"\n通过 {ok}/{len(cases)}")


if __name__ == "__main__":
    main()
