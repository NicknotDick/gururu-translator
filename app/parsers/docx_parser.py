"""Word (.docx) 解析器：段落级翻译，保留样式、编号与表格结构。"""
from __future__ import annotations

import copy
import re
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

from ..core.ir import is_translatable
from .base import MODE_BILINGUAL, BaseParser, Holder

# 目录（TOC）是自动生成的域，不应被当作正文翻译；翻译后需标记脏位让 Word 重算。
_TOC_RE = re.compile(r"\bTOC\b", re.IGNORECASE)
# 手动（非域）目录的标题，用于启用静态目录条目识别。
_TOC_HEADING_RE = re.compile(r"(table of contents|contents|^目录$|^目\s*录$)", re.IGNORECASE)
# 静态目录条目：以编号开头、引导点结尾、末尾为页码数字。
_TOC_ENTRY_RE = re.compile(r"^\s*\d+(?:\.\d+)*\.?\s+\S.*\.{2,}\s*\d{1,4}\s*$", re.UNICODE)


def _has_fldchar(p_el: Any, char_type: str) -> bool:
    for fc in p_el.iter(qn("w:fldChar")):
        if fc.get(qn("w:fldCharType")) == char_type:
            return True
    return False


def _para_has_toc_fldsimple(p_el: Any) -> bool:
    for fs in p_el.iter(qn("w:fldSimple")):
        instr = fs.get(qn("w:instr")) or ""
        if _TOC_RE.search(instr):
            return True
    return False


def _para_has_toc_instr(p_el: Any) -> bool:
    for it in p_el.iter(qn("w:instrText")):
        if _TOC_RE.search(it.text or ""):
            return True
    return False


def _doc_has_toc_heading(body: Any) -> bool:
    """文档是否含目录标题（目录 / Table of Contents 等），用于启用静态目录识别。"""
    for p_el in body.iter(qn("w:p")):
        if _TOC_HEADING_RE.search(_para_text(p_el).strip()):
            return True
    return False


def _para_is_static_toc_entry(p_el: Any) -> bool:
    """判断段落是否为手动（非域）目录条目：编号 + 引导点 + 末尾页码。"""
    text = _para_text(p_el).strip()
    if not text:
        return False
    return bool(_TOC_ENTRY_RE.search(text))


def _iter_body_with_toc_skip(body: Any, static_ok: bool = False):
    """按顺序遍历正文段落，并标记属于目录（TOC）域的段落（应跳过）。

    目录域有两种写法：
    - w:fldSimple（自包含）：整段跳过即可。
    - 复杂域（begin/separate/end 跨多个 w:p）：从含 TOC 指令的段落起，
      直到 end 域字符所在段落，全部跳过。
    """
    in_toc = False
    for p_el in body.iter(qn("w:p")):
        if in_toc:
            if _has_fldchar(p_el, "end"):
                in_toc = False
            yield p_el, True
            continue
        if _para_has_toc_fldsimple(p_el):
            yield p_el, True
            continue
        if _para_has_toc_instr(p_el):
            in_toc = not _has_fldchar(p_el, "end")
            yield p_el, True
            continue
        if static_ok and _para_is_static_toc_entry(p_el):
            yield p_el, True
            continue
        yield p_el, False


def _mark_toc_dirty(doc: Any) -> None:
    """给目录域打脏位，Word/WPS 打开时自动重算（修正错乱/不靠右的页码）。"""
    body = doc.element.body
    for fs in body.iter(qn("w:fldSimple")):
        instr = fs.get(qn("w:instr")) or ""
        if _TOC_RE.search(instr):
            fs.set(qn("w:dirty"), "true")
    for p_el in body.iter(qn("w:p")):
        if _para_has_toc_instr(p_el):
            for fc in p_el.iter(qn("w:fldChar")):
                if fc.get(qn("w:fldCharType")) == "begin":
                    fc.set(qn("w:dirty"), "true")


def _ensure_update_fields(doc: Any) -> None:
    """在 settings.xml 写入 w:updateFields，强制 Word/WPS 打开时重算全部域
    （含目录/页码），确保目录页码右对齐且正确。"""
    try:
        settings = doc.settings.element
    except Exception:
        return
    if settings is None:
        return
    for child in settings.iter(qn("w:updateFields")):
        child.set(qn("w:val"), "true")
        return
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)


def _in_textbox(el: Any, root: Any) -> bool:
    anc = el.getparent()
    while anc is not None and anc is not root:
        if anc.tag == qn("w:txbxContent"):
            return True
        anc = anc.getparent()
    return False


def _para_text(p_el: Any) -> str:
    parts: list[str] = []
    for t in p_el.iter(qn("w:t")):
        if _in_textbox(t, p_el):
            continue
        parts.append(t.text or "")
    return "".join(parts)


def _set_para_text(p_el: Any, text: str, parent: Any) -> None:
    ts = [t for t in p_el.iter(qn("w:t")) if not _in_textbox(t, p_el)]
    if not ts:
        Paragraph(p_el, parent).add_run(text)
        return
    ts[0].text = text
    ts[0].set(qn("xml:space"), "preserve")
    for t in ts[1:]:
        t.text = ""


def _strip_numbering(p_el: Any) -> None:
    pPr = p_el.find(qn("w:pPr"))
    if pPr is None:
        return
    numPr = pPr.find(qn("w:numPr"))
    if numPr is not None:
        pPr.remove(numPr)


def _drop_drawings(p_el: Any) -> None:
    for tag in ("w:drawing", "w:pict"):
        for el in list(p_el.iter(qn(tag))):
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)


class DocxParser(BaseParser):
    ext = ".docx"
    label = "Word 文档"
    supports_bilingual = True

    # ---------- 扫描 ----------
    def _open_and_scan(self, path: str, opts: dict) -> tuple[Any, list[Holder]]:
        doc = Document(path)
        parent = doc._body
        holders: list[Holder] = []
        body_el = doc.element.body
        # 仅当文档含目录标题时，才识别并跳过手动（非域）目录条目，
        # 避免误伤正文中的普通段落。
        static_ok = _doc_has_toc_heading(body_el)

        def add(p_el: Any, kind: str = "para") -> None:
            text = _para_text(p_el)

            def _set(value: str, _el=p_el) -> None:
                _set_para_text(_el, value, parent)

            holders.append(Holder(kind=kind, text=text, setter=_set, obj=p_el))

        for p_el, skip in _iter_body_with_toc_skip(body_el, static_ok=static_ok):
            if skip:
                continue
            add(p_el)

        if not opts.get("skip_header_footer"):
            seen: set[int] = set()
            for section in doc.sections:
                for part in (section.header, section.footer,
                             section.first_page_header, section.first_page_footer,
                             section.even_page_header, section.even_page_footer):
                    try:
                        el = part._element
                    except Exception:
                        continue
                    if el is None or id(el) in seen:
                        continue
                    seen.add(id(el))
                    for p_el in el.iter(qn("w:p")):
                        if _para_has_toc_fldsimple(p_el):
                            continue
                        add(p_el, "meta")

        return doc, holders

    def scan(self, path: str, opts: dict) -> list[Holder]:
        return self._open_and_scan(path, opts)[1]

    # ---------- 回写 ----------
    def commit(self, path: str, out_path: str, mapping: dict[int, str],
               mode: str, opts: dict) -> dict:
        doc, holders = self._open_and_scan(path, opts)
        overflow: list[str] = []
        for i, h in enumerate(holders):
            tgt = mapping.get(i)
            if not tgt:
                continue
            if mode == MODE_BILINGUAL:
                src_el = h.obj
                if len(tgt) > 900:
                    overflow.append(f"第 {i + 1} 段译文过长，可能溢出")
                try:
                    new_el = copy.deepcopy(src_el)
                    _drop_drawings(new_el)
                    _strip_numbering(new_el)
                    _set_para_text(new_el, tgt, doc._body)
                    src_el.addnext(new_el)
                except Exception:
                    if h.setter:
                        h.setter(tgt)
            else:
                try:
                    if h.setter:
                        h.setter(tgt)
                except Exception:
                    pass
        # 目录域标记脏位 + 强制更新域，打开时由 Word/WPS 重算页码
        # （避免错乱 / 不靠右），双保险确保目录页码右对齐且正确。
        _mark_toc_dirty(doc)
        _ensure_update_fields(doc)
        doc.save(out_path)
        return {"overflow": overflow}

    def preview_meta(self, path: str) -> dict:
        try:
            doc = Document(path)
            return {"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}
        except Exception:
            return {}
