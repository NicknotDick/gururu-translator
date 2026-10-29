"""Excel (.xlsx) 解析器：单元格级翻译，双语模式在右侧插入译文列。"""
from __future__ import annotations

import re
import zipfile
from copy import copy
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..core.ir import is_translatable
from .base import MODE_BILINGUAL, BaseParser, Holder

CELL_REF = re.compile(r"(?<![A-Za-z0-9_$])(\$?)([A-Z]{1,3})(\$?)(\d{1,7})(?![A-Za-z0-9_])")
QUOTED = re.compile(r'"[^"]*"')


def _col_of(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def _shift_formula(formula: str, source_cols: list[int]) -> str:
    """插列后平移公式中的列引用。"""
    if not formula.startswith("=") or not source_cols:
        return formula
    pieces = QUOTED.split(formula)
    quotes = QUOTED.findall(formula)

    def repl(m: re.Match) -> str:
        col = _col_of(m.group(2))
        delta = sum(1 for s in source_cols if s < col)
        if not delta:
            return m.group(0)
        return f"{m.group(1)}{get_column_letter(col + delta)}{m.group(3)}{m.group(4)}"

    out: list[str] = []
    for i, piece in enumerate(pieces):
        out.append(CELL_REF.sub(repl, piece))
        if i < len(quotes):
            out.append(quotes[i])
    return "".join(out)


def _is_formula(v: Any) -> bool:
    return isinstance(v, str) and v.startswith("=")


class XlsxParser(BaseParser):
    ext = ".xlsx"
    label = "Excel 工作簿"
    supports_bilingual = True

    # ---------- 打开与扫描 ----------
    def _open_and_scan(self, path: str, opts: dict):
        wb = load_workbook(path, data_only=False, rich_text=False)
        parent = wb._body if hasattr(wb, "_body") else None  # noqa: F841
        holders: list[Holder] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if not isinstance(v, str) or _is_formula(v):
                        continue
                    if not is_translatable(v):
                        continue
                    holders.append(Holder(
                        kind="cell",
                        text=v,
                        obj=cell,
                        extra={"sheet": ws.title, "row": cell.row, "col": cell.column,
                               "group": ws.title},
                    ))
        return wb, holders

    def scan(self, path: str, opts: dict) -> list[Holder]:
        return self._open_and_scan(path, opts)[1]

    def load(self, path, opts=None):
        ir = super().load(path, opts)
        try:
            with zipfile.ZipFile(path) as z:
                names = z.namelist()
            if any(n.startswith("xl/media/") for n in names):
                ir.warnings.append("工作簿含内嵌图片，重新保存后可能丢失")
            if any(n.startswith("xl/charts/") for n in names):
                ir.warnings.append("工作簿含图表，重新保存后图表样式可能变化")
        except Exception:
            pass
        return ir

    # ---------- 回写 ----------
    def commit(self, path: str, out_path: str, mapping: dict[int, str],
               mode: str, opts: dict) -> dict:
        wb, holders = self._open_and_scan(path, opts)
        warnings: list[str] = []

        if mode != MODE_BILINGUAL:
            for i, h in enumerate(holders):
                tgt = mapping.get(i)
                if not tgt:
                    continue
                ex = h.extra
                try:
                    wb[ex["sheet"]].cell(row=ex["row"], column=ex["col"]).value = tgt
                except Exception:
                    continue
            wb.save(out_path)
            return {"warnings": warnings}

        # 按工作表分组
        by_sheet: dict[str, dict[tuple[int, int], int]] = {}
        for i, h in enumerate(holders):
            tgt = mapping.get(i)
            if not tgt:
                continue
            ex = h.extra
            by_sheet.setdefault(ex["sheet"], {})[(ex["row"], ex["col"])] = i
        if not by_sheet:
            wb.save(out_path)
            return {"warnings": warnings}

        for sheet_name, cells in by_sheet.items():
            ws = wb[sheet_name]
            asc = sorted({c for (_r, c) in cells})
            if len(asc) > 60:
                warnings.append(f"{sheet_name}：需插入 {len(asc)} 列，表格会明显变宽")

            widths = {d.min: d.width for d in ws.column_dimensions.values() if d.width}
            merged = [(r.min_col, r.max_col, r.min_row, r.max_row) for r in ws.merged_cells.ranges]
            for r in list(ws.merged_cells.ranges):
                try:
                    ws.unmerge_cells(str(r))
                except Exception:
                    pass

            for c in sorted(asc, reverse=True):
                ws.insert_cols(c + 1)

            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and v.startswith("="):
                        try:
                            cell.value = _shift_formula(v, asc)
                        except Exception:
                            pass

            for (row, col), holder_idx in cells.items():
                extra = sum(1 for s in asc if s < col)
                src_cell = ws.cell(row=row, column=col + extra)
                new_cell = ws.cell(row=row, column=col + extra + 1)
                try:
                    new_cell._style = copy(src_cell._style)
                except Exception:
                    pass
                new_cell.value = mapping[holder_idx]

            for (min_c, max_c, min_r, max_r) in merged:
                nm = min_c + sum(1 for s in asc if s < min_c)
                nx = max_c + sum(1 for s in asc if s <= max_c)
                try:
                    ws.merge_cells(
                        start_row=min_r, end_row=max_r,
                        start_column=nm, end_column=nx,
                    )
                except Exception:
                    pass

            try:
                for key in list(ws.column_dimensions.keys()):
                    del ws.column_dimensions[key]
                for c, w in widths.items():
                    extra = sum(1 for s in asc if s < c)
                    ws.column_dimensions[get_column_letter(c + extra)].width = w
                for c in asc:
                    extra = sum(1 for s in asc if s < c)
                    ws.column_dimensions[get_column_letter(c + extra + 1)].width = widths.get(c, 10)
            except Exception:
                pass

            if ws.freeze_panes:
                warnings.append(f"{sheet_name}：冻结窗格未随插列调整")

        wb.save(out_path)
        return {"warnings": warnings}
